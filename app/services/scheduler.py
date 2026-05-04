import asyncio
import logging
import threading
import uuid
import json
import tempfile
from io import BytesIO
from pathlib import Path
from collections import Counter
from typing import Callable

from sqlalchemy.exc import IntegrityError
from sqlalchemy import func, select
from openpyxl import Workbook

from datetime import date, datetime, UTC
from datetime import timedelta

from app.core.config import settings
from app.db.database import SessionLocal
from app.db.models import (
    Document,
    DocumentStatus,
    EmailGroup,
    NotificationSettings,
    PipelineRun,
    PipelineRunFile,
    PipelineRunStatus,
)
from app.services.auth.microsoft_auth import MicrosoftAuthError, MicrosoftAuthService
from app.services.email_service import EmailServiceError, send_email
from app.services.email_templates import build_eod_summary_email, build_pipeline_email
from app.services.orchestrator import Orchestrator, OrchestratorError
from app.services.storage.onedrive import (
    INTAKE_FOLDER,
    PROCESSED_FOLDER,
    OUTPUT_SUCCESS_FOLDER,
    OUTPUT_FAILURE_FOLDER,
    OneDriveClient,
    OneDriveError,
    OneDriveNotFoundError,
    ensure_pipeline_folders,
)

logger = logging.getLogger(__name__)

INTAKE_DIR = Path("storage/intake")
IGNORED_SUFFIXES = {".tmp", ".part", ".swp"}
ALLOWED_INPUT_EXTENSIONS = {".pdf", ".doc", ".docx"}
DENIED_EXTENSIONS = {".xlsx", ".json", ".html"}
GENERATED_ARTIFACT_SUFFIXES = (
    ".converted.pdf",
    ".tagged.pdf",
    "_tagged_pdf.pdf",
    "_accessibility_report.json",
    "_tagged_report.xlsx",
    ".autotag-report.xlsx",
)


class Scheduler:
    def __init__(self, interval: int | None = None, status_listener: Callable[[dict[str, object]], None] | None = None) -> None:
        configured_interval = interval or settings.SCHEDULER_INTERVAL_SECONDS
        self.interval = max(1, configured_interval)
        self.max_files_per_cycle = max(1, settings.MAX_FILES_PER_CYCLE)
        self.max_cycles_per_run = max(1, settings.MAX_CYCLES_PER_RUN)
        self.processing_delay_seconds = max(0, settings.PROCESSING_DELAY_SECONDS)
        self.storage_provider = settings.STORAGE_PROVIDER
        self.automation_enabled = True
        self.auth_service = MicrosoftAuthService() if self.storage_provider == "onedrive" else None
        self.onedrive_client = OneDriveClient() if self.storage_provider == "onedrive" else None
        self._thread: threading.Thread | None = None
        self._eod_thread: threading.Thread | None = None
        self._stop_event = threading.Event()
        self._wakeup_event = threading.Event()
        self.last_failure: dict | None = None
        self._status_listener = status_listener
        self._status_lock = threading.Lock()
        self._onedrive_folders_ensured = False
        self.pipeline_status: dict[str, object] = {
            "is_running": False,
            "current_step": None,
            "current_file": None,
            "progress": 0,
            "error": None,
            "failed_step": None,
        }

    def _record_failure(self, *, filename: str, failed_step: str, error: str, progress: int) -> None:
        self.last_failure = {
            "at": datetime.now(UTC),
            "current_file": filename,
            "failed_step": failed_step,
            "error": error,
            "progress": progress,
        }
        self._set_pipeline_status(
            current_step="FAILED",
            current_file=filename,
            progress=int(max(0, min(100, progress))),
            error=error,
            failed_step=failed_step,
        )

    def _set_pipeline_status(self, **updates: object) -> None:
        snapshot: dict[str, object]
        with self._status_lock:
            self.pipeline_status.update(updates)
            snapshot = dict(self.pipeline_status)
        if self._status_listener:
            try:
                self._status_listener(snapshot)
            except Exception:
                logger.exception("Failed to publish pipeline status update")

    def _mark_pipeline_started(self) -> None:
        self._set_pipeline_status(
            is_running=True,
            current_step="RUNNING",
            current_file=None,
            progress=0,
            error=None,
            failed_step=None,
        )

    def _mark_pipeline_progress(self, *, filename: str, progress: int) -> None:
        self._set_pipeline_status(
            is_running=True,
            current_step="RUNNING",
            current_file=filename,
            progress=int(max(0, min(99, progress))),
            error=None,
            failed_step=None,
        )

    def _mark_pipeline_finished(self, *, success: bool) -> None:
        self._set_pipeline_status(
            is_running=False,
            current_step="COMPLETED" if success else "FAILED",
            current_file=None,
            progress=0,
        )

    def get_pipeline_status(self) -> dict[str, object]:
        with self._status_lock:
            return dict(self.pipeline_status)

    def start(self) -> None:
        if self._thread and self._thread.is_alive():
            logger.info("Scheduler already running")
            return

        self._stop_event.clear()
        self._thread = threading.Thread(target=self._run_loop, name="document-scheduler", daemon=True)
        self._eod_thread = threading.Thread(target=self._run_eod_loop, name="eod-scheduler", daemon=True)
        self._thread.start()
        self._eod_thread.start()
        logger.info(
            "Scheduler started: interval=%s seconds storage_provider=%s",
            self.interval,
            self.storage_provider,
        )

    def stop(self) -> None:
        self._stop_event.set()
        if self._thread and self._thread.is_alive():
            self._thread.join(timeout=5)
        if self._eod_thread and self._eod_thread.is_alive():
            self._eod_thread.join(timeout=5)
        logger.info("Scheduler stopped")

    def status(self) -> str:
        if self._thread and self._thread.is_alive():
            return "running"
        return "stopped"

    def update_interval(self, new_interval: int) -> None:
        self.interval = max(1, new_interval)
        logger.info("Scheduler interval updated: interval=%s seconds", self.interval)
        # Wake the scheduler loop so the new interval takes effect immediately.
        self._wakeup_event.set()

    async def run_once(self) -> None:
        self._mark_pipeline_started()
        run_id = str(uuid.uuid4())
        started_at = datetime.now(UTC)
        logger.info("Manual scheduler run started: run_id=%s", run_id)
        run_db_id = await self._create_pipeline_run(run_id=run_id, start_time=started_at)

        files_summary: list[dict[str, object]] = []
        access_passed = 0
        access_failed = 0
        access_manual = 0
        total_processed = 0
        total_failed = 0
        total_skipped = 0
        run_failed = False

        try:
            run_result = await self._drain_intake_for_run()
            files_summary = run_result.get("files", [])
            access = run_result.get("accessibility", {})
            access_passed = int(access.get("passed", 0))
            access_failed = int(access.get("failed", 0))
            access_manual = int(access.get("manual", 0))
            total_processed = int(run_result.get("processed", 0))
            total_failed = int(run_result.get("failed", 0))
            total_skipped = int(run_result.get("skipped", 0))
        except Exception:
            run_failed = True
            raise
        finally:
            finished_at = datetime.now(UTC)
            duration = self._format_duration(finished_at - started_at)
            total_files = len(files_summary)
            success_count = sum(1 for item in files_summary if item.get("status") == DocumentStatus.COMPLETED.value)
            failure_count = sum(1 for item in files_summary if item.get("status") == DocumentStatus.FAILED.value)
            final_status = PipelineRunStatus.FAILED if run_failed else PipelineRunStatus.COMPLETED

            if run_db_id:
                await self._store_pipeline_run_files(run_db_id=run_db_id, files_summary=files_summary)
                await self._finalize_pipeline_run(
                    run_db_id=run_db_id,
                    end_time=finished_at,
                    duration=duration,
                    total_files=total_files,
                    success_count=success_count,
                    failure_count=failure_count,
                    status=final_status,
                )

            run_data = {
                "run_id": run_id,
                "duration": duration,
                "total_files": total_files,
                "success_count": success_count,
                "failure_count": failure_count,
                "total_processed": total_processed,
                "total_failed": total_failed,
                "total_skipped": total_skipped,
                "files": files_summary,
                "passed": access_passed,
                "failed": access_failed,
                "manual": access_manual,
                "dashboard_url": settings.FRONTEND_DASHBOARD_URL,
            }

            try:
                recipient_emails = await self._get_notification_emails()
                if not recipient_emails:
                    logger.info("Skipping email notification (no configured recipients): run_id=%s", run_id)
                else:
                    if failure_count == 0:
                        subject = f"✅ Pipeline Completed | {total_files} Files"
                    else:
                        subject = f"❌ Pipeline Completed with Errors | {failure_count} Failed"
                    html_content = build_pipeline_email(run_data)
                    for recipient in recipient_emails:
                        await send_email(to_email=recipient, subject=subject, html_content=html_content)
                    logger.info(
                        "Pipeline summary email sent: run_id=%s recipients=%s",
                        run_id,
                        ",".join(recipient_emails),
                    )
            except EmailServiceError:
                logger.exception("Email notification failed: run_id=%s", run_id)
            except Exception:
                logger.exception("Unexpected error while sending email: run_id=%s", run_id)

            self._mark_pipeline_finished(success=not run_failed)
            logger.info("Manual scheduler run finished: run_id=%s duration=%s", run_id, duration)

    def _run_loop(self) -> None:
        while not self._stop_event.is_set():
            signaled = self._wakeup_event.wait(self.interval)
            if self._stop_event.is_set():
                break
            if signaled:
                # Wakeups are used to apply scheduler config changes immediately,
                # not to force an immediate processing run.
                self._wakeup_event.clear()
                continue
            logger.info("Scheduler polling cycle started")
            try:
                asyncio.run(self.run_once())
            except Exception:
                logger.exception("Scheduler polling cycle failed")
            logger.info("Scheduler polling cycle ended")

    def _run_eod_loop(self) -> None:
        while not self._stop_event.is_set():
            try:
                self._maybe_send_eod_summary()
            except Exception:
                logger.exception("EOD summary check failed")
            self._stop_event.wait(60)

    def _poll_local_intake_folder(self) -> dict:
        candidate_files: list[Path] = []
        for path in sorted(INTAKE_DIR.iterdir(), key=lambda item: item.name.lower()):
            is_processible, reason = self._is_processible_file(path)
            if is_processible:
                candidate_files.append(path)
                continue

            if reason == "generated":
                logger.info("Scheduler skipped generated file: filename=%s", path.name)
            elif reason == "invalid":
                logger.info("Scheduler skipped invalid file: filename=%s", path.name)

        process_batch = candidate_files[: self.max_files_per_cycle]
        logger.info(
            "Scheduler batch prepared: candidates=%d processing_now=%d deferred=%d max_per_cycle=%d",
            len(candidate_files),
            len(process_batch),
            max(0, len(candidate_files) - len(process_batch)),
            self.max_files_per_cycle,
        )

        processed_count = 0
        skipped_count = len(candidate_files) - len(process_batch)
        failed_count = 0
        files_summary: list[dict[str, object]] = []
        access_passed = 0
        access_failed = 0
        access_manual = 0

        for index, file_path in enumerate(process_batch):
            total_files = max(1, len(process_batch))
            self._mark_pipeline_progress(
                filename=file_path.name,
                progress=int((index / total_files) * 100),
            )
            result = self._register_and_trigger(file_path=file_path, source_id=None)
            if result:
                processed_count += 1
                rp = str(result.get("report_path") or "")
                summary = result.get("accessibility") or {"passed": 0, "failed": 0, "manual": 0}
                status = str(result.get("status", ""))
                if status == DocumentStatus.FAILED.value:
                    failed_count += 1
                files_summary.append(
                    {
                        "name": result.get("name", ""),
                        "status": status,
                        "error": result.get("error", ""),
                        "output_stem": Scheduler._output_stem_from_report_path(rp),
                        "accessibility": summary,
                    }
                )
                access_passed += int(summary.get("passed", 0))
                access_failed += int(summary.get("failed", 0))
                access_manual += int(summary.get("manual", 0))
            else:
                skipped_count += 1
            if index < len(process_batch) - 1 and self.processing_delay_seconds > 0:
                self._stop_event.wait(self.processing_delay_seconds)

        return {
            "processed": processed_count,
            "skipped": skipped_count,
            "failed": failed_count,
            "files": files_summary,
            "accessibility": {"passed": access_passed, "failed": access_failed, "manual": access_manual},
        }

    async def _drain_intake_for_run(self) -> dict[str, object]:
        files_summary: list[dict[str, object]] = []
        total_processed = 0
        total_failed = 0
        total_skipped = 0
        access_passed = 0
        access_failed = 0
        access_manual = 0
        cycle_number = 0

        while cycle_number < self.max_cycles_per_run:
            cycle_number += 1
            if self.storage_provider == "onedrive":
                cycle_result = await self.process_onedrive_intake()
            elif self.storage_provider == "local":
                cycle_result = self._poll_local_intake_folder()
            else:
                raise RuntimeError(
                    f"Unsupported storage provider '{self.storage_provider}'. "
                    "Expected one of: onedrive, local."
                )

            processed_this_cycle = int(cycle_result.get("processed", 0))
            failed_this_cycle = int(cycle_result.get("failed", 0))
            skipped_this_cycle = int(cycle_result.get("skipped", 0))
            cycle_access = cycle_result.get("accessibility", {})

            total_processed += processed_this_cycle
            total_failed += failed_this_cycle
            total_skipped += skipped_this_cycle
            access_passed += int(cycle_access.get("passed", 0))
            access_failed += int(cycle_access.get("failed", 0))
            access_manual += int(cycle_access.get("manual", 0))
            files_summary.extend(cycle_result.get("files", []))

            if processed_this_cycle == 0:
                logger.info(
                    "Scheduler drain cycle stopped: cycle=%d processed_this_cycle=%d reason=no_files_processed",
                    cycle_number,
                    processed_this_cycle,
                )
                break

            logger.info(
                "Scheduler drain cycle complete: cycle=%d processed_this_cycle=%d",
                cycle_number,
                processed_this_cycle,
            )

            if cycle_number >= self.max_cycles_per_run:
                logger.info(
                    "Scheduler drain cycle stopped: cycle=%d processed_this_cycle=%d reason=max_cycles_reached max_cycles_per_run=%d",
                    cycle_number,
                    processed_this_cycle,
                    self.max_cycles_per_run,
                )
                break

        return {
            "files": files_summary,
            "processed": total_processed,
            "failed": total_failed,
            "skipped": total_skipped,
            "cycle_count": cycle_number,
            "accessibility": {"passed": access_passed, "failed": access_failed, "manual": access_manual},
        }

    async def process_onedrive_intake(self) -> dict:
        if not self.onedrive_client:
            raise OneDriveError("OneDrive client is not initialized.")
        if not self.auth_service:
            raise MicrosoftAuthError("Microsoft auth service is not initialized.")

        if not self._onedrive_folders_ensured:
            try:
                await ensure_pipeline_folders(await self._get_onedrive_access_token())
                self._onedrive_folders_ensured = True
            except Exception:
                logger.exception("Failed to ensure required OneDrive pipeline folders before intake processing.")

        try:
            all_files = await self.onedrive_client.list_files(
                access_token=await self._get_onedrive_access_token(),
                folder_path=INTAKE_FOLDER,
            )
        except (OneDriveError, MicrosoftAuthError):
            logger.exception("Scheduler failed to list OneDrive intake folder: folder=%s", INTAKE_FOLDER)
            raise

        candidate_files: list[dict[str, str]] = []
        for remote_file in sorted(all_files, key=lambda item: item["name"].lower()):
            filename = Path(remote_file["name"]).name
            is_processible, reason = self._is_processible_filename(filename)
            if is_processible:
                candidate_files.append(remote_file)
                continue
            if reason == "generated":
                logger.info("Scheduler skipped generated OneDrive file: filename=%s file_id=%s", filename, remote_file["id"])
            elif reason == "invalid":
                logger.info("Scheduler skipped invalid OneDrive file: filename=%s file_id=%s", filename, remote_file["id"])

        # Apply max_files_per_cycle to files that are not already known in DB,
        # so previously processed files do not consume the current run capacity.
        async with SessionLocal() as db:
            candidate_source_ids = [item["id"] for item in candidate_files]
            existing_source_ids: set[str] = set()
            if candidate_source_ids:
                existing_rows = (
                    await db.execute(select(Document.filename).where(Document.filename.in_(candidate_source_ids)))
                )
                existing_rows = existing_rows.all()
                existing_source_ids = {row[0] for row in existing_rows}

        process_batch = []
        for remote_file in candidate_files:
            if remote_file["id"] in existing_source_ids:
                continue
            process_batch.append(remote_file)
            if len(process_batch) >= self.max_files_per_cycle:
                break
        logger.info(
            "OneDrive intake batch prepared: candidates=%d unprocessed=%d processing_now=%d deferred=%d max_per_cycle=%d",
            len(candidate_files),
            max(0, len(candidate_files) - len(existing_source_ids)),
            len(process_batch),
            max(0, len(candidate_files) - len(existing_source_ids) - len(process_batch)),
            self.max_files_per_cycle,
        )

        processed_count = 0
        skipped_count = len(candidate_files) - len(process_batch)
        failed_count = 0
        files_summary: list[dict[str, object]] = []
        access_passed = 0
        access_failed = 0
        access_manual = 0
        for index, remote_file in enumerate(process_batch):
            file_id = remote_file["id"]
            filename = Path(remote_file["name"]).name
            total_files = max(1, len(process_batch))
            self._mark_pipeline_progress(
                filename=filename,
                progress=int((index / total_files) * 100),
            )
            try:
                with tempfile.TemporaryDirectory(prefix="onedrive-intake-") as temp_dir:
                    local_path = Path(temp_dir) / f"{file_id}_{filename}"
                    await self.onedrive_client.download_file(
                        access_token=await self._get_onedrive_access_token(),
                        file_id=file_id,
                        local_path=local_path,
                    )
                    logger.info("File fetched from OneDrive: file_id=%s filename=%s local_path=%s", file_id, filename, local_path)
                    result = await self._register_and_trigger_onedrive(
                        file_path=local_path,
                        source_id=file_id,
                        original_filename=filename,
                    )
                if result:
                    processed_count += 1
                    rp = str(result.get("report_path") or "")
                    summary = result.get("accessibility") or {"passed": 0, "failed": 0, "manual": 0}
                    files_summary.append(
                        {
                            "name": result.get("name", ""),
                            "status": result.get("status", ""),
                            "error": result.get("error", ""),
                            "output_stem": Scheduler._output_stem_from_report_path(rp),
                            "accessibility": summary,
                        }
                    )
                    access_passed += int(summary.get("passed", 0))
                    access_failed += int(summary.get("failed", 0))
                    access_manual += int(summary.get("manual", 0))
                else:
                    skipped_count += 1
            except OneDriveNotFoundError:
                failed_count += 1
                logger.error("OneDrive file no longer exists: file_id=%s filename=%s", file_id, filename)
                files_summary.append(
                    {
                        "name": filename,
                        "status": DocumentStatus.FAILED.value,
                        "error": "OneDrive file no longer exists",
                        "output_stem": "",
                    }
                )
            except OneDriveError:
                failed_count += 1
                logger.exception("Failed to download/process OneDrive file: file_id=%s filename=%s", file_id, filename)
                files_summary.append(
                    {
                        "name": filename,
                        "status": DocumentStatus.FAILED.value,
                        "error": "Failed to download/process OneDrive file",
                        "output_stem": "",
                    }
                )
            if index < len(process_batch) - 1 and self.processing_delay_seconds > 0:
                self._stop_event.wait(self.processing_delay_seconds)

        return {
            "processed": processed_count,
            "skipped": skipped_count,
            "failed": failed_count,
            "files": files_summary,
            "accessibility": {"passed": access_passed, "failed": access_failed, "manual": access_manual},
        }

    async def _register_and_trigger(
        self,
        file_path: Path,
        source_id: str | None,
        original_filename: str | None = None,
    ) -> dict[str, str] | None:
        display_filename = original_filename or file_path.name
        filename = file_path.name
        logger.info("Scheduler evaluating file: filename=%s source_id=%s", display_filename, source_id or "n/a")
        async with SessionLocal() as db:
            try:
                dedupe_key = source_id or filename
                existing = (await db.execute(select(Document).where(Document.filename == dedupe_key))).scalars().first()
                if existing:
                    if existing.status == DocumentStatus.FAILED:
                        logger.info(
                            "Scheduler skipping recently failed file: filename=%s source_id=%s document_id=%s",
                            display_filename,
                            source_id or "n/a",
                            existing.id,
                        )
                    else:
                        logger.info(
                            "Scheduler skipping known file: filename=%s source_id=%s document_id=%s",
                            display_filename,
                            source_id or "n/a",
                            existing.id,
                        )
                    return {
                        "name": display_filename,
                        "status": existing.status.value,
                        "error": "Already processed.",
                        "report_path": "",
                    }

                document = Document(filename=dedupe_key, status=DocumentStatus.UPLOADED)
                db.add(document)
                await db.commit()
                await db.refresh(document)
                logger.info(
                    "Scheduler created document: document_id=%s dedupe_key=%s source_file=%s",
                    document.id,
                    dedupe_key,
                    display_filename,
                )

                orchestrator = Orchestrator(db=db)
                result = await orchestrator.process_document(document_id=document.id, input_path=file_path)

                if self.storage_provider == "onedrive":
                    try:
                        await self._upload_onedrive_outputs(
                            tagged_pdf_path=result["tagged_pdf_path"],
                            report_path=result["report_path"],
                            autotag_report_path=result.get("autotag_report_path"),
                        )
                    except OneDriveError:
                        document.status = DocumentStatus.FAILED
                        db.add(document)
                        await db.commit()
                        logger.exception("Scheduler failed to upload outputs to OneDrive: document_id=%s", document.id)
                        return {
                            "name": display_filename,
                            "status": DocumentStatus.FAILED.value,
                            "error": "Failed to upload outputs to OneDrive.",
                            "report_path": result.get("report_path", ""),
                            "accessibility": {"passed": 0, "failed": 0, "manual": 0},
                        }
                report_path = str(result.get("report_path") or "")
                summary = self._read_accessibility_summary(report_path) if report_path else {"passed": 0, "failed": 0, "manual": 0}
                logger.info("Scheduler triggered processing: document_id=%s", document.id)
                return {
                    "name": display_filename,
                    "status": DocumentStatus.COMPLETED.value,
                    "error": "",
                    "report_path": report_path,
                    "accessibility": summary,
                }
            except IntegrityError:
                await db.rollback()
                logger.info("Scheduler encountered duplicate file insert: filename=%s", display_filename)
                return {"name": display_filename, "status": "SKIPPED", "error": "Duplicate insert ignored.", "report_path": ""}
            except OrchestratorError as exc:
                logger.exception("Scheduler processing failed for file: filename=%s", display_filename)
                failed_step = getattr(exc, "failed_step", "") or "unknown"
                progress_map = {"fetch": 5, "convert_to_pdf": 30, "precheck": 45, "auto_tag": 60, "check_accessibility": 80}
                self._record_failure(
                    filename=display_filename,
                    failed_step=failed_step,
                    error=str(exc) or "Pipeline failed.",
                    progress=int(progress_map.get(failed_step, 40)),
                )
                return {
                    "name": display_filename,
                    "status": DocumentStatus.FAILED.value,
                    "error": str(exc) or "Pipeline failed.",
                    "report_path": "",
                    "accessibility": {"passed": 0, "failed": 0, "manual": 0},
                }
            except Exception:
                logger.exception("Scheduler failed while handling file: filename=%s", display_filename)
                return {
                    "name": display_filename,
                    "status": DocumentStatus.FAILED.value,
                    "error": "Scheduler failed while handling file.",
                    "report_path": "",
                    "accessibility": {"passed": 0, "failed": 0, "manual": 0},
                }

    async def _register_and_trigger_onedrive(
        self,
        file_path: Path,
        source_id: str,
        original_filename: str,
    ) -> dict[str, str] | None:
        document: Document | None = None
        async with SessionLocal() as db:
            try:
                existing = (await db.execute(select(Document).where(Document.filename == source_id))).scalars().first()
                if existing:
                    logger.info(
                        "Skipping already processed OneDrive file: source_id=%s filename=%s document_id=%s",
                        source_id,
                        original_filename,
                        existing.id,
                    )
                    return {
                        "name": original_filename,
                        "status": existing.status.value,
                        "error": "Already processed.",
                        "report_path": "",
                    }

                document = Document(filename=source_id, status=DocumentStatus.UPLOADED)
                db.add(document)
                await db.commit()
                await db.refresh(document)
                logger.info(
                    "Created document for OneDrive file: document_id=%s source_id=%s filename=%s",
                    document.id,
                    source_id,
                    original_filename,
                )

                orchestrator = Orchestrator(db=db)
                result = await orchestrator.process_document(document_id=document.id, input_path=file_path)
                await self.onedrive_client.move_file(
                    access_token=await self._get_onedrive_access_token(),
                    file_id=source_id,
                    folder_path=PROCESSED_FOLDER,
                    filename=original_filename,
                )
                logger.info("OneDrive original moved to processed: file_id=%s filename=%s", source_id, original_filename)
                await self._upload_onedrive_outputs(
                    tagged_pdf_path=result["tagged_pdf_path"],
                    report_path=result["report_path"],
                    autotag_report_path=result.get("autotag_report_path"),
                )
                report_path = str(result.get("report_path") or "")
                summary = self._read_accessibility_summary(report_path) if report_path else {"passed": 0, "failed": 0, "manual": 0}
                logger.info("Processed OneDrive file successfully: source_id=%s document_id=%s", source_id, document.id)
                return {
                    "name": original_filename,
                    "status": DocumentStatus.COMPLETED.value,
                    "error": "",
                    "report_path": report_path,
                    "accessibility": summary,
                }
            except IntegrityError:
                await db.rollback()
                logger.info("Duplicate OneDrive document insert ignored: source_id=%s", source_id)
                return {"name": original_filename, "status": "SKIPPED", "error": "Duplicate insert ignored.", "report_path": ""}
            except OrchestratorError as exc:
                logger.exception("Processing failed for OneDrive file: source_id=%s filename=%s", source_id, original_filename)
                try:
                    await self.onedrive_client.move_file(
                        access_token=await self._get_onedrive_access_token(),
                        file_id=source_id,
                        folder_path=OUTPUT_FAILURE_FOLDER,
                        filename=original_filename,
                    )
                    logger.info("OneDrive original moved to failure: file_id=%s filename=%s", source_id, original_filename)
                except OneDriveError:
                    logger.exception("Failed to move OneDrive original to failure: file_id=%s", source_id)
                failed_step = getattr(exc, "failed_step", "") or "unknown"
                progress_map = {"fetch": 5, "convert_to_pdf": 30, "precheck": 45, "auto_tag": 60, "check_accessibility": 80}
                self._record_failure(
                    filename=original_filename,
                    failed_step=failed_step,
                    error=str(exc) or "Pipeline failed.",
                    progress=int(progress_map.get(failed_step, 40)),
                )
                return {
                    "name": original_filename,
                    "status": DocumentStatus.FAILED.value,
                    "error": str(exc) or "Pipeline failed.",
                    "report_path": "",
                    "accessibility": {"passed": 0, "failed": 0, "manual": 0},
                }
            except OneDriveError:
                if document:
                    document.status = DocumentStatus.FAILED
                    db.add(document)
                    await db.commit()
                logger.exception("Upload failed for OneDrive outputs: source_id=%s document_id=%s", source_id, document.id if document else "n/a")
                return {
                    "name": original_filename,
                    "status": DocumentStatus.FAILED.value,
                    "error": "Upload failed for OneDrive outputs.",
                    "report_path": "",
                    "accessibility": {"passed": 0, "failed": 0, "manual": 0},
                }

    @staticmethod
    def _format_duration(delta: timedelta) -> str:
        seconds = int(max(0, delta.total_seconds()))
        if seconds < 60:
            return f"{seconds}s"
        minutes, sec = divmod(seconds, 60)
        if minutes < 60:
            return f"{minutes}m {sec}s"
        hours, minutes = divmod(minutes, 60)
        return f"{hours}h {minutes}m {sec}s"

    @staticmethod
    def _output_stem_from_report_path(report_path: str) -> str:
        """Stem used for *_accessibility_report.json / OneDrive uploads (may differ from display file_name)."""
        if not report_path:
            return ""
        name = Path(report_path).name
        lower = name.lower()
        for suffix in ("_accessibility_report.json", "_report.json"):
            if lower.endswith(suffix.lower()):
                return name[: -len(suffix)]
        if lower.endswith(".accessibility-report.json"):
            return name[: -len(".accessibility-report.json")]
        return ""

    @staticmethod
    def _read_accessibility_summary(report_path: str) -> dict[str, int]:
        try:
            path = Path(report_path)
            if not path.exists():
                return {"passed": 0, "failed": 0, "manual": 0}
            raw = path.read_text(encoding="utf-8", errors="replace")
            payload = json.loads(raw)
            summary = payload.get("Summary") or {}
            passed = int(summary.get("Passed", 0)) + int(summary.get("Passed manually", 0))
            failed = int(summary.get("Failed", 0)) + int(summary.get("Failed manually", 0))
            manual = int(summary.get("Needs manual check", 0))
            return {"passed": passed, "failed": failed, "manual": manual}
        except Exception:
            logger.exception("Failed to parse accessibility report summary: path=%s", report_path)
            return {"passed": 0, "failed": 0, "manual": 0}

    @staticmethod
    async def _get_notification_emails() -> list[str]:
        async with SessionLocal() as db:
            try:
                rows = (await db.execute(select(EmailGroup.email).order_by(EmailGroup.created_at.asc()))).all()
                emails = [row[0].strip().lower() for row in rows if row and row[0]]
                deduped = list(dict.fromkeys(emails))
                return deduped
            except Exception:
                logger.exception("Failed to resolve notification email list")
                return []

    @staticmethod
    async def _create_pipeline_run(run_id: str, start_time: datetime) -> str:
        async with SessionLocal() as db:
            try:
                row = PipelineRun(run_id=run_id, start_time=start_time, status=PipelineRunStatus.COMPLETED)
                db.add(row)
                await db.commit()
                await db.refresh(row)
                return row.id
            except Exception:
                await db.rollback()
                logger.exception("Failed to create pipeline run row: run_id=%s", run_id)
                return ""

    @staticmethod
    async def _store_pipeline_run_files(run_db_id: str, files_summary: list[dict[str, object]]) -> None:
        if not run_db_id or not files_summary:
            return
        async with SessionLocal() as db:
            try:
                for item in files_summary:
                    file_status = item.get("status")
                    status = PipelineRunStatus.COMPLETED if file_status == DocumentStatus.COMPLETED.value else PipelineRunStatus.FAILED
                    stem_val = str(item.get("output_stem") or "").strip() or None
                    accessibility = item.get("accessibility") or {}
                    db.add(
                        PipelineRunFile(
                            run_id=run_db_id,
                            file_name=str(item.get("name") or ""),
                            output_stem=stem_val,
                            status=status,
                            error_message=str(item.get("error") or ""),
                            accessibility_passed=int(accessibility.get("passed", 0)),
                            accessibility_failed=int(accessibility.get("failed", 0)),
                            accessibility_manual=int(accessibility.get("manual", 0)),
                        )
                    )
                await db.commit()
            except Exception:
                await db.rollback()
                logger.exception("Failed to store pipeline run file rows: run_db_id=%s", run_db_id)

    @staticmethod
    async def _finalize_pipeline_run(
        run_db_id: str,
        end_time: datetime,
        duration: str,
        total_files: int,
        success_count: int,
        failure_count: int,
        status: PipelineRunStatus,
    ) -> None:
        if not run_db_id:
            return
        async with SessionLocal() as db:
            try:
                row = (await db.execute(select(PipelineRun).where(PipelineRun.id == run_db_id))).scalars().first()
                if not row:
                    return
                row.end_time = end_time
                row.duration = duration
                row.total_files = int(total_files)
                row.success_count = int(success_count)
                row.failure_count = int(failure_count)
                row.status = status
                db.add(row)
                await db.commit()
            except Exception:
                await db.rollback()
                logger.exception("Failed to finalize pipeline run row: run_db_id=%s", run_db_id)

    def _maybe_send_eod_summary(self) -> None:
        asyncio.run(self._maybe_send_eod_summary_async())

    async def _maybe_send_eod_summary_async(self) -> None:
        async with SessionLocal() as db:
            try:
                settings_row = (
                    await db.execute(select(NotificationSettings).order_by(NotificationSettings.created_at.asc()))
                ).scalars().first()
                if not settings_row:
                    settings_row = NotificationSettings(eod_time="18:00", enabled=False)
                    db.add(settings_row)
                    await db.commit()
                    await db.refresh(settings_row)

                if not settings_row.enabled:
                    return

                now = datetime.now()
                current_hhmm = now.strftime("%H:%M")
                if current_hhmm != settings_row.eod_time:
                    return

                today = now.date()
                if settings_row.last_sent_date == today:
                    return

                day_start = datetime.combine(today, datetime.min.time())
                day_end = day_start + timedelta(days=1)
                runs = (
                    await db.execute(
                        select(PipelineRun)
                        .where(PipelineRun.created_at >= day_start, PipelineRun.created_at < day_end)
                        .order_by(PipelineRun.start_time.asc())
                    )
                ).scalars().all()
                total_runs = len(runs)
                total_files = sum(int(item.total_files or 0) for item in runs)
                total_success = sum(int(item.success_count or 0) for item in runs)
                total_failure = sum(int(item.failure_count or 0) for item in runs)

                run_rows = [
                    {
                        "run_id": item.run_id,
                        "start_time": item.start_time,
                        "total_files": item.total_files,
                        "success_count": item.success_count,
                        "failure_count": item.failure_count,
                        "duration": item.duration,
                        "status": item.status.value,
                    }
                    for item in runs
                ]

                failure_messages = (
                    await db.execute(
                        select(PipelineRunFile.error_message, func.count(PipelineRunFile.id))
                        .join(PipelineRun, PipelineRun.id == PipelineRunFile.run_id)
                        .where(
                            PipelineRun.created_at >= day_start,
                            PipelineRun.created_at < day_end,
                            PipelineRunFile.status == PipelineRunStatus.FAILED,
                            PipelineRunFile.error_message != "",
                        )
                        .group_by(PipelineRunFile.error_message)
                    )
                ).all()
                if failure_messages:
                    common_error = Counter({msg: count for msg, count in failure_messages}).most_common(1)[0][0]
                else:
                    common_error = ""

                payload = {
                    "date": today,
                    "runs": run_rows,
                    "totals": {"runs": total_runs, "files": total_files, "success": total_success, "failure": total_failure},
                    "common_error": common_error,
                }
                subject = f"📊 Daily Pipeline Summary | {today.isoformat()}"
                html_content = build_eod_summary_email(payload)
                attachment_name = f"daily_pipeline_summary_{today.isoformat()}.xlsx"
                attachment_bytes = self._build_eod_summary_xlsx(
                    report_date=today,
                    runs=run_rows,
                    totals={
                        "runs": total_runs,
                        "files": total_files,
                        "success": total_success,
                        "failure": total_failure,
                    },
                )
                recipients = await self._get_notification_emails()
                if not recipients:
                    logger.info("Skipping EOD summary email: no recipients")
                else:
                    for recipient in recipients:
                        await send_email(
                            to_email=recipient,
                            subject=subject,
                            html_content=html_content,
                            attachments=[
                                (
                                    attachment_name,
                                    attachment_bytes,
                                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                                )
                            ],
                        )
                    logger.info("EOD summary sent: date=%s recipients=%s", today.isoformat(), ",".join(recipients))

                settings_row.last_sent_date = today
                db.add(settings_row)
                await db.commit()
            except EmailServiceError:
                await db.rollback()
                logger.exception("Failed to send EOD summary email")
            except Exception:
                await db.rollback()
                logger.exception("Failed during EOD summary processing")

    @staticmethod
    def _build_eod_summary_xlsx(
        report_date: date,
        runs: list[dict[str, object]],
        totals: dict[str, int],
    ) -> bytes:
        workbook = Workbook()
        runs_sheet = workbook.active
        runs_sheet.title = "Runs"
        runs_sheet.append(["Run ID", "Time", "Duration", "Files", "Success", "Failed", "Status"])
        for run in runs:
            start_time_raw = run.get("start_time")
            if isinstance(start_time_raw, datetime):
                start_time_value = start_time_raw.isoformat()
            else:
                start_time_value = str(start_time_raw or "")
            runs_sheet.append(
                [
                    str(run.get("run_id") or ""),
                    start_time_value,
                    str(run.get("duration") or "-"),
                    int(run.get("total_files") or 0),
                    int(run.get("success_count") or 0),
                    int(run.get("failure_count") or 0),
                    str(run.get("status") or ""),
                ]
            )

        output = BytesIO()
        workbook.save(output)
        return output.getvalue()

    async def _upload_onedrive_outputs(
        self,
        tagged_pdf_path: str,
        report_path: str,
        autotag_report_path: str | None = None,
    ) -> None:
        if not self.onedrive_client:
            raise OneDriveError("OneDrive client is not initialized for upload.")

        tagged_path_obj = Path(tagged_pdf_path)
        report_path_obj = Path(report_path)
        await self.onedrive_client.upload_file(
            access_token=await self._get_onedrive_access_token(),
            local_path=tagged_path_obj,
            folder_path=OUTPUT_SUCCESS_FOLDER,
            filename=tagged_path_obj.name,
        )
        logger.info("Upload success to OneDrive: local=%s target_folder=%s", tagged_path_obj, OUTPUT_SUCCESS_FOLDER)
        await self.onedrive_client.upload_file(
            access_token=await self._get_onedrive_access_token(),
            local_path=report_path_obj,
            folder_path=OUTPUT_SUCCESS_FOLDER,
            filename=report_path_obj.name,
        )
        logger.info("Upload success to OneDrive: local=%s target_folder=%s", report_path_obj, OUTPUT_SUCCESS_FOLDER)
        if autotag_report_path:
            autotag_report_obj = Path(autotag_report_path)
            if autotag_report_obj.exists():
                await self.onedrive_client.upload_file(
                    access_token=await self._get_onedrive_access_token(),
                    local_path=autotag_report_obj,
                    folder_path=OUTPUT_SUCCESS_FOLDER,
                    filename=autotag_report_obj.name,
                )
                logger.info("Upload success to OneDrive: local=%s target_folder=%s", autotag_report_obj, OUTPUT_SUCCESS_FOLDER)

    async def _get_onedrive_access_token(self) -> str:
        if not self.auth_service:
            raise MicrosoftAuthError("Microsoft auth service is not initialized.")
        async with SessionLocal() as db:
            return await self.auth_service.get_valid_access_token(db=db)

    @staticmethod
    def _is_processible_file(path: Path) -> tuple[bool, str]:
        if not path.is_file():
            return False, "ignore"
        return Scheduler._is_processible_filename(path.name)

    @staticmethod
    def _is_processible_filename(name: str) -> tuple[bool, str]:
        if name.startswith(".") or name.startswith("~"):
            return False, "ignore"
        if Path(name).suffix.lower() in IGNORED_SUFFIXES:
            return False, "ignore"

        lower_name = name.lower()
        if any(lower_name.endswith(suffix) for suffix in GENERATED_ARTIFACT_SUFFIXES):
            return False, "generated"

        suffix = Path(name).suffix.lower()
        if suffix in DENIED_EXTENSIONS:
            return False, "invalid"
        if suffix not in ALLOWED_INPUT_EXTENSIONS:
            return False, "invalid"
        return True, "valid"

